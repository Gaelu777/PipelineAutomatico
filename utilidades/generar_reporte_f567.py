"""Genera el reporte Word de las Fases 5, 6 y 7, con el mismo formato que los anteriores."""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

RAIZ = Path(__file__).resolve().parent.parent
F4 = RAIZ / "salidas_fase4"
F6 = RAIZ / "salidas_fase6"
AZUL = RGBColor(0x1F, 0x3A, 0x5F)

doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = doc.sections[0].left_margin


def titulo_bloque(texto, size, bold=True, color=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(texto)
    r.font.size = Pt(size)
    r.bold = bold
    if color:
        r.font.color.rgb = color
    return p


def H(texto, nivel=1):
    return doc.add_heading(texto, level=nivel)


def P(texto, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
    p.add_run(texto)
    p.paragraph_format.space_after = Pt(6)
    return p


def B(texto):
    return doc.add_paragraph(texto, style="List Bullet")


def tabla(headers, filas, nota=None):
    t = doc.add_table(rows=len(filas) + 1, cols=len(headers))
    t.style = "Table Grid"
    for j, h in enumerate(headers):
        c = t.rows[0].cells[j]
        c.text = str(h)
        for par in c.paragraphs:
            for r in par.runs:
                r.bold = True
                r.font.size = Pt(8)
    for i, fila in enumerate(filas):
        for j, v in enumerate(fila):
            c = t.rows[i + 1].cells[j]
            c.text = str(v)
            for par in c.paragraphs:
                for r in par.runs:
                    r.font.size = Pt(8)
    if nota:
        p = doc.add_paragraph()
        r = p.add_run(nota)
        r.font.size = Pt(8)
        r.italic = True
    doc.add_paragraph()
    return t


# ------------------------------------------------------------------ datos
esc = json.loads((F6 / "parametros_escalamiento.json").read_text(encoding="utf-8"))
res = json.loads((F6 / "resumen_fase5_fase6.json").read_text(encoding="utf-8"))
seg = pd.read_csv(F6 / "segmentos_fase5.csv")
cuar = pd.read_csv(F6 / "resumen_cuarentena.csv")
estres = pd.read_csv(F6 / "pruebas_estres.csv")
enganos = pd.read_csv(F4 / "datos_enganosos.csv")
limpias = pd.read_csv(F6 / "transacciones_limpias.csv")

# ----------------------------------------------------------------- portada
titulo_bloque("Reporte de Transformación, Validación y Carga: Fases 5, 6 y 7", 20, True, AZUL)
titulo_bloque("Proyecto Integrador: Pipeline ETL Robusto con Zona de Cuarentena", 11)
titulo_bloque("Ingeniería de Datos (ATD-2602) | Unidad 2 | TecNM Campus Colima | Semestre Ago-Dic 2026", 9, False)
titulo_bloque(
    "Alcance de este reporte: Fase 5 (escalamiento Min-Max y discretización), Fase 6 (validación "
    "Pydantic y zona de cuarentena) y Fase 7 (carga transaccional a SQLite). Grupo 1: Retail "
    "omnicanal (tienda física + tienda web).", 9, False)

# --------------------------------------------------------------- 1. resumen
H("1. Resumen ejecutivo")
P(f"Se tomó la salida de la Fase 4 ({res['entrada']:,} registros ya limpios, integrados y con huérfanos "
  f"marcados) y se completó el pipeline hasta la base de datos. La Fase 5 normalizó el importe al rango "
  f"[0, 1] y lo segmentó en tres niveles; la Fase 6 lo validó contra un esquema estricto de Pydantic; la "
  f"Fase 7 lo persistió en SQLite con transacciones seguras.")
P(f"De {res['entrada']:,} registros, {res['validos']:,} ({res['tasa_aceptacion_pct']}%) pasaron el esquema "
  f"y {res['cuarentena']:,} ({round(100 - res['tasa_aceptacion_pct'], 2)}%) se desviaron a cuarentena con su "
  f"motivo documentado. Ninguna fila se perdió: {res['validos']:,} + {res['cuarentena']:,} = {res['entrada']:,}.")

tabla(["Etapa", "Registros", "Porcentaje"],
      [["Entrada (salida Fase 4)", f"{res['entrada']:,}", "100.00%"],
       ["transacciones_limpias", f"{res['validos']:,}", f"{res['tasa_aceptacion_pct']}%"],
       ["cuarentena_errores", f"{res['cuarentena']:,}", f"{round(100 - res['tasa_aceptacion_pct'], 2)}%"]])

P("El resultado que importa para la defensa no es la tasa de aceptación, sino que el pipeline procesó "
  "los 1,189 registros sin detenerse una sola vez. Las 24 pruebas de estrés con datos deliberadamente "
  "corruptos terminaron con cero excepciones no controladas.")

# ------------------------------------------------------- 2. datos engañosos
H("2. Diagnóstico previo: datos que parecen correctos pero no lo son")
P("Las Fases 1 a 4 atacan errores visibles: nulos, duplicados exactos, formatos rotos, outliers. Antes de "
  "escribir el esquema de validación se hizo un análisis forense distinto, buscando los registros que pasan "
  "cualquier revisión de tipo y formato y aun así son falsos. Son los que ensucian un reporte sin que nadie "
  "lo note, porque en un vistazo rápido se ven perfectamente normales.")
P("Script: utilidades/forense_datos_enganosos.py. Se detectaron diez tipos de error silencioso:")

tabla(["Cód.", "Hallazgo", "Afectados", "Por qué engaña"],
      [[r["codigo"], r["titulo"], int(r["afectados"]), r["por_que_enganan"][:115]]
       for _, r in enganos.iterrows()],
      "La columna 'Afectados' se mide sobre la fuente donde aplica cada regla, por eso los totales no suman.")

P("Dos de estos hallazgos cambiaron el diseño del validador:")
B("E-09 obligó a que la regla aritmética respete flag_outlier. Si la Fase 3 recortó el monto por IQR, la fila "
  "deja de cumplir monto = cantidad × precio por culpa de nuestra propia transformación, no del sistema origen. "
  "Mandarla a cuarentena sería castigar un dato que nosotros rompimos.")
B("E-10 no se manda a cuarentena. 119 registros con estatus CANCELADO conservan importe; descartar el 10% del "
  "lote sería destructivo y un pedido cancelado sigue siendo un hecho real. Se resuelve con la columna derivada "
  "monto_facturable_mxn, que vale 0 cuando el estatus es CANCELADO o DEVUELTO. El dato se conserva; lo que se "
  "corrige es cuánto factura.")

P("El caso más ilustrativo es E-03. Los 29 huérfanos de tienda física cotizan todos a 500.00 exacto. Un precio "
  "de 500 pesos es perfectamente plausible y pasa cualquier validación de rango; lo que lo delata es la "
  "repetición. Un precio real no cae 29 veces en el mismo número redondo: es el valor de relleno que el punto "
  "de venta escribe cuando no encuentra el producto en el catálogo.")

# ------------------------------------------------------------- 3. fase 5
H("3. Fase 5: Escalamiento Min-Max y discretización")

H("3.1 Escalamiento Min-Max", 2)
P("Se normalizaron monto_mxn y cantidad al rango [0, 1] con la fórmula (x − mín) / (máx − mín).")
mm, mc = esc["minmax_monto"], esc["minmax_cantidad"]
tabla(["Variable", "Mínimo", "Máximo", "Rango resultante", "Fuera de [0,1]"],
      [["monto_mxn", f"{mm['min']:,.2f}", f"{mm['max']:,.2f}", "[0, 1]", "0"],
       ["cantidad", f"{mc['min']:,.0f}", f"{mc['max']:,.0f}", "[0, 1]", "0"]])

P("Min-Max depende únicamente del mínimo y el máximo, lo que la convierte en la técnica de escalamiento más "
  "frágil ante valores extremos. Por eso se aplica después del capping de la Fase 3 y no antes: con el monto "
  "sin recortar el máximo era 2,109,055 pesos, y ese único valor habría comprimido más del 99% de las ventas "
  "por debajo de 0.01, dejando la variable normalizada inservible. Con los extremos ya tratados, el rango "
  "[332.67, 16,510.02] reparte la escala entre los datos reales.", "Por qué después del capping: ")

H("3.2 Discretización en tres segmentos", 2)
P("El importe se partió en BAJO, MEDIO y ALTO con pd.qcut(q=3), es decir por terciles de la distribución.")
tabla(["Segmento", "Registros", "Mínimo (MXN)", "Máximo (MXN)", "Promedio (MXN)"],
      [[r["segmento_monto"], int(r["n"]), f"{r['minimo']:,.2f}", f"{r['maximo']:,.2f}",
        f"{r['promedio']:,.2f}"] for _, r in seg.iterrows()])

P("El catálogo mezcla categorías con precios muy distintos (Alimentos frente a "
  "Electrónica). Un corte fijo en pesos (por ejemplo 0–1000, 1000–5000, +5000) metería una categoría completa "
  "dentro del mismo segmento y podría dejar grupos casi vacíos. Los terciles garantizan tres grupos de tamaño "
  "comparable (403, 391 y 395 registros), de modo que cualquier estadístico calculado por segmento tiene un "
  "respaldo muestral similar.", "Por qué terciles y no cortes fijos: ")
P(f"Los puntos de corte quedaron en {esc['terciles']['cortes']} y se guardan en "
  f"parametros_escalamiento.json, de modo que la segmentación es reproducible y auditable en cualquier "
  f"ejecución posterior.")

# ------------------------------------------------------------- 4. fase 6
H("4. Fase 6: Validación con Pydantic y zona de cuarentena")

H("4.1 Esquema y reglas", 2)
P("Se definió la clase TransaccionLimpia(BaseModel) con validadores por campo (@field_validator) y "
  "validadores de fila completa (@model_validator(mode='after')) para las reglas que ningún campo puede "
  "comprobar por sí solo. Cada regla cruzada responde a un hallazgo del diagnóstico forense.")

tabla(["Regla", "Qué exige", "Hallazgo que la origina"],
      [["R1", "La llave de producto no puede ser huérfana", "E-04"],
       ["R2", "precio_unitario_mxn distinto de 500.00 exacto", "E-03"],
       ["R3", "monto = cantidad × precio (± $0.01), salvo si flag_outlier", "E-01, E-02, E-09"],
       ["R4", "Estatus no facturable implica monto_facturable_mxn = 0", "E-10"],
       ["R5", "Toda venta ONLINE exige correo de contacto", "E-05"],
       ["Campo", "id_producto cumple ^PROD-\\d{3}$", "E-04"],
       ["Campo", "Correo con regex anclado; rechaza '@@' y '..'", "E-05"],
       ["Campo", "fecha dentro de [2026-01-01, 2026-12-31]", "E-07"],
       ["Campo", "estatus dentro del catálogo de seis valores", "—"]])

H("4.2 Bifurcación y motivos de rechazo", 2)
P("Cada registro se valida aislado dentro de un bloque try/except ValidationError. Una fila corrupta se "
  "desvía con su motivo y el lote continúa: en ningún caso una fila mala detiene el procesamiento de las "
  "demás, que es el principio de robustez que exige el proyecto.")

tabla(["Causa de rechazo", "Registros", "% del lote"],
      [[r["causa"], int(r["registros"]), f"{100*r['registros']/res['entrada']:.2f}%"]
       for _, r in cuar.iterrows()]
      + [["Total en cuarentena", res["cuarentena"],
          f"{round(100 - res['tasa_aceptacion_pct'], 2)}%"]])

P("Los 60 huérfanos del archivo unificado están los 60 en cuarentena, pero solo 57 "
  "aparecen etiquetados como tales. Los otros 3 figuran bajo su error de correo porque Pydantic evalúa los "
  "validadores de campo antes que el validador de fila, de modo que la etiqueta refleja el primer error "
  "encontrado y no el único. La columna motivo_rechazo_cuarentena conserva la lista completa de violaciones "
  "de cada registro. Verificación directa sobre la tabla final: 0 huérfanos se filtraron a "
  "transacciones_limpias.", "Nota de lectura: ")

H("4.3 Pruebas de estrés", 2)
P(f"Se inyectaron {len(estres)} registros deliberadamente corruptos. El criterio de éxito no es que el "
  f"pipeline los acepte, sino que los rechace con un mensaje explícito y sin lanzar una excepción no "
  f"controlada.")
rech = int((estres["resultado"] == "RECHAZADO").sum())
acep = int((estres["resultado"] == "ACEPTADO").sum())
crash = int((estres["resultado"] == "CRASH").sum())
tabla(["Resultado", "Casos", "Interpretación"],
      [["RECHAZADO", rech, "El validador detectó el problema y lo explicó"],
       ["ACEPTADO", acep, "Únicamente el caso de control, que es válido"],
       ["CRASH (excepción no controlada)", crash, "Debe ser 0 para aprobar resiliencia"]])

P("Muestra de los casos más representativos:")
muestra = estres[estres["caso"].isin([
    "texto en campo numerico", "fecha imposible 31/02", "llave huerfana",
    "precio de relleno 500", "total inflado x80 (bot)", "correo con doble arroba",
    "nulo en toda la fila", "caso valido de control"])]
tabla(["Caso inyectado", "Resultado", "Mensaje del validador"],
      [[r["caso"], r["resultado"], str(r["detalle"])[:78]] for _, r in muestra.iterrows()])

# ------------------------------------------------------------- 5. fase 7
H("5. Fase 7: Carga transaccional a base de datos")

H("5.1 Esquema e integridad declarada", 2)
P("Se persistieron las dos tablas obligatorias en SQLite. La integridad no se confía solo al código Python: "
  "queda declarada en el propio esquema de la base, de modo que un INSERT incorrecto es rechazado por el "
  "motor aunque la capa de aplicación fallara.")
B("PRIMARY KEY en id_registro, que impide cargar dos veces la misma venta.")
B("CHECK en cantidad > 0, precios > 0 y monto_mxn_norm entre 0 y 1.")
B("CHECK de dominio en canal (FISICA/ONLINE) y segmento_monto (BAJO/MEDIO/ALTO).")
B("GLOB 'PROD-[0-9][0-9][0-9]' sobre id_producto, el patrón de llave válido.")
B("NOT NULL en los campos que el negocio exige e índices por fecha, canal y producto.")

tabla(["Tabla", "Filas cargadas", "Contenido"],
      [["transacciones_limpias", f"{res['validos']:,}", "Registros que superaron el esquema Pydantic"],
       ["cuarentena_errores", f"{res['cuarentena']:,}", "Registros desviados, con causa y motivo detallado"],
       ["Total", f"{res['entrada']:,}", "Cuadra exactamente con la entrada de la Fase 5"]])

H("5.2 Transacciones y prueba de rollback", 2)
P("Cada tabla se carga dentro de una única transacción: BEGIN, DELETE previo (que hace la carga idempotente), "
  "executemany e commit al final. Cualquier excepción dispara rollback y deja la tabla exactamente como estaba.")
P("No basta con afirmar que el rollback funciona: se probó contra la base real. Se "
  "intentó insertar un registro con cantidad = −99, que viola el CHECK (cantidad > 0). El motor lanzó "
  "IntegrityError, el pipeline lo capturó, ejecutó rollback y el conteo de la tabla antes y después de la "
  "operación fue idéntico: 1,057 filas. La atomicidad queda demostrada, no asumida.",
  "Prueba explícita: ")

H("5.3 Resultado de negocio", 2)
total_fact = limpias["monto_facturable_mxn"].sum()
piv = (limpias.groupby(["canal", "segmento_monto"])
       .agg(pedidos=("id_registro", "count"), facturado=("monto_facturable_mxn", "sum"))
       .reset_index())
orden = {"BAJO": 0, "MEDIO": 1, "ALTO": 2}
piv = piv.sort_values(["canal", "segmento_monto"], key=lambda s: s.map(orden) if s.name == "segmento_monto" else s)
tabla(["Canal", "Segmento", "Pedidos", "Facturado (MXN)"],
      [[r["canal"], r["segmento_monto"], int(r["pedidos"]), f"{r['facturado']:,.2f}"]
       for _, r in piv.iterrows()]
      + [["TOTAL", "—", int(len(limpias)), f"{total_fact:,.2f}"]])

fis = piv[piv["canal"] == "FISICA"]["facturado"].sum()
P(f"La tienda física concentra el {100*fis/total_fact:.1f}% del importe facturado con menos "
  f"pedidos que el canal online. Su segmento ALTO por sí solo (265 pedidos) factura más que el canal online "
  f"completo. El online vende volumen en el segmento BAJO: 202 pedidos que apenas representan "
  f"{100*piv[(piv.canal=='ONLINE')&(piv.segmento_monto=='BAJO')]['facturado'].sum()/total_fact:.1f}% del "
  f"importe. Son dos negocios con perfiles distintos, y la discretización de la Fase 5 es lo que permite "
  f"verlo de un vistazo.", "Lectura: ")

# ------------------------------------------------------------ 6. conclusion
H("6. Conclusiones")
P(f"El pipeline quedó completo de extremo a extremo: de {res['entrada']:,} registros crudos integrados a "
  f"{res['validos']:,} ventas validadas en base de datos, con {res['cuarentena']:,} registros aislados en "
  f"cuarentena y cada uno con su causa documentada. Ninguna fila se perdió ni se duplicó en el trayecto.")
P("La lección técnica del tramo final es que la validación de tipos es la parte fácil. Lo que realmente "
  "contamina una base son los registros que respetan todos los tipos: el total inflado exactamente ochenta "
  "veces, el precio de relleno de 500 pesos, la llave con formato impecable que apunta a un producto "
  "inexistente, el correo con doble arroba que una comprobación ingenua da por bueno. Esos errores solo "
  "aparecen al cruzar campos entre sí y contra el catálogo maestro, y por eso el peso del esquema recae en "
  "el validador de fila completa y no en los validadores de campo.")
P("La decisión de diseño más defendible del proyecto es la de E-10: no todo registro inconsistente merece "
  "cuarentena. Descartar los 119 pedidos cancelados habría sido técnicamente correcto y comercialmente "
  "absurdo. Separar el hecho (la venta ocurrió y luego se canceló) de su efecto contable (no factura) "
  "mediante una columna derivada conserva la información y corrige el reporte. Una zona de cuarentena bien "
  "diseñada aísla lo que no se puede confiar, no todo lo que no encaja.")

doc.add_paragraph()
P("Scripts: fase5_fase6_pipeline.py, fase7_carga_bd.py y utilidades/forense_datos_enganosos.py. "
  "Reglas y cifras completas en REGLAS_DE_NEGOCIO.md. Repositorio: github.com/Gaelu777/PipelineAutomatico")

salida = RAIZ / "Reporte_Fase5_Fase6_Fase7_Validacion_y_Carga.docx"
doc.save(salida)
print(f"OK -> {salida}")
print(f"   parrafos: {len(doc.paragraphs)}  tablas: {len(doc.tables)}")
